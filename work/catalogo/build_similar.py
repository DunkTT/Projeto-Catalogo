from pathlib import Path
import json,math,unicodedata
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT,WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
root=Path(__file__).parent
out=Path('C:/Users/estoque/Documents/Codex/2026-10-08/erifique-se-o-mcp-do-bling/outputs/catalogo-produtos-ilustrado.docx')
products=json.loads((root/'products.json').read_text(encoding='utf8'))
images={x['id']:x for x in json.loads((root/'similar/final_images.json').read_text(encoding='utf8'))}
key=lambda p:unicodedata.normalize('NFD',p['nome']).encode('ascii','ignore').decode().upper()
products=sorted(products,key=key)
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.55)
sec.left_margin=sec.right_margin=Inches(.55);sec.header_distance=sec.footer_distance=Inches(.25)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(3)
 for e in list(st.element.xpath('.//w:pBdr')):e.getparent().remove(e)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.0
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(12)
doc.core_properties.title='Catálogo ilustrado de produtos'
doc.core_properties.subject='Produtos por nome e SKU com fotografias de modelos iguais ou semelhantes'
doc.core_properties.author=''
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('Catálogo ilustrado  |  ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
def borders(table):
 b=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
 table._tbl.tblPr.append(b)
def pad(cell):
 m=OxmlElement('w:tcMar')
 for side in ['top','left','bottom','right']:
  e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
 cell._tc.get_or_add_tcPr().append(m)
def para(cell,text,size=9,bold=False):
 p=cell.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3)
 for run in p.runs:run.font.size=Pt(size);run.bold=bold
 return p
def link(p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');z=OxmlElement('w:sz');z.set(qn('w:val'),'15');pr.append(z)
 col=OxmlElement('w:color');col.set(qn('w:val'),'555555');pr.append(col);r.append(pr)
 t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
visited=[]
for page in range(math.ceil(len(products)/9)):
 if page:
  doc.add_page_break()
  doc.add_heading('Catálogo de produtos',level=1)
 else:
  doc.add_paragraph('Catálogo ilustrado de produtos',style='Title')
  p=doc.add_paragraph('1.693 itens • Nomes e códigos SKU • Outubro de 2026')
  for run in p.runs:run.font.size=Pt(10)
  p=doc.add_paragraph('As fotos ilustrativas mostram modelos semelhantes à descrição. Aparência, cor e especificações podem variar. Os nomes e SKUs foram preservados.')
  for run in p.runs:run.font.size=Pt(9)
  p.paragraph_format.space_after=Pt(8)
 group=products[page*9:page*9+9]
 table=doc.add_table(rows=math.ceil(len(group)/3),cols=3);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c in table.columns:c.width=Inches(7.4/3)
 borders(table)
 for idx,item in enumerate(group):
  row=table.rows[idx//3];row.height=Inches(2.65 if page==0 else 2.8);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  cell=row.cells[idx%3];cell.width=Inches(7.4/3);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pad(cell)
  p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5)
  im=images.get(item['id'])
  if im:
   w,h=Image.open(im['file']).size;scale=min(2.12/w,1.4/h)
   shape=p.add_run().add_picture(im['file'],width=Inches(w*scale),height=Inches(h*scale));shape._inline.docPr.set('descr',item['nome']+(' — foto ilustrativa de modelo semelhante' if im['approx'] else ''))
  else:
   text='Serviço sem foto' if item['id']=='16650946480' else 'Imagem pendente de identificação'
   p.add_run(text).font.size=Pt(10)
   p.paragraph_format.space_before=Pt(32);p.paragraph_format.space_after=Pt(32)
  p=para(cell,item['nome'],9.5,True)
  n=len(visited)+1;e=OxmlElement('w:bookmarkStart');e.set(qn('w:id'),str(n));e.set(qn('w:name'),'produto_'+item['id']);p._p.insert(0,e)
  e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(n));p._p.append(e);visited.append(item['id'])
  para(cell,'SKU '+(item['sku'] or 'não cadastrado'),9)
  if im:
   p=para(cell,'Foto ilustrativa — modelo semelhante' if im['approx'] else ('Foto do cadastro' if im['source']=='Bling' else 'Foto do modelo identificado'),7.5)
   for run in p.runs:run.font.color.rgb=RGBColor.from_string('555555')
   p=para(cell,'',7.5);p.paragraph_format.space_after=Pt(0)
   if str(im['source']).startswith('http'):link(p,'Fonte da imagem',im['source'])
  elif item['id']!='16650946480':para(cell,'Descrição insuficiente para selecionar uma foto',7.5)
 spacer=doc.add_paragraph();spacer.paragraph_format.line_spacing=Pt(1);spacer.paragraph_format.space_after=Pt(0);spacer.add_run().font.size=Pt(1)
assert len(set(visited))==len(products)==1693
doc.save(out)
meta={'output':str(out),'total':1693,'photos':len(images),'approx':sum(x['approx'] for x in images.values()),'pagesPlanned':math.ceil(len(products)/9),'bytes':out.stat().st_size}
(root/'similar/build.json').write_text(json.dumps(meta),encoding='utf8')
print(json.dumps(meta,ensure_ascii=False))

